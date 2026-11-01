"""Render three standalone audit documents from PLATFORM-SCALE-AUDIT.md."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
SRC = (ROOT / "PLATFORM-SCALE-AUDIT.md").read_text(encoding="utf-8").splitlines()

NAVY = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x0F, 0x6E, 0x6E)
BODY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
HEADER_BG = "1F4E79"
ZEBRA = "F4F7FB"


def set_run_font(run, name="Calibri", size=11, bold=False, color=BODY, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "E5E7EB")
        borders.append(el)
    tc_pr.append(borders)


def add_inline(paragraph, text, size=11, color=BODY, base_bold=False):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, bold=base_bold, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(size - 1, 9), color=NAVY)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, bold=base_bold, color=color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def style_heading(paragraph, level):
    sizes = {1: 20, 2: 16, 3: 13}
    colors = {1: NAVY, 2: TEAL, 3: NAVY}
    before = {1: 18, 2: 14, 3: 10}
    for run in paragraph.runs:
        set_run_font(run, size=sizes[level], bold=True, color=colors[level])
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before[level])
    pf.space_after = Pt(6)
    pf.line_spacing = 1.08
    pf.keep_with_next = True


def body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    add_inline(p, text)
    return p


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j in range(cols):
            cell = row.cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            value = row_data[j] if j < len(row_data) else ""
            white = RGBColor(0xFF, 0xFF, 0xFF)
            add_inline(
                p,
                value.strip(),
                size=9,
                color=white if i == 0 else BODY,
                base_bold=(i == 0),
            )
            shade_cell(cell, HEADER_BG if i == 0 else (ZEBRA if i % 2 == 0 else "FFFFFF"))
            set_cell_border(cell)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def is_separator(line):
    stripped = line.replace("|", "").replace(":", "").replace("-", "").strip()
    return stripped == "" and "-" in line


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def render(lines, out_path, eyebrow, title, subtitle, footer_label):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(footer_label)
    set_run_font(run, size=9, bold=True, color=NAVY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("1 million / day  ·  page ")
    set_run_font(label, size=9, color=MUTED)
    add_page_number(footer)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY

    cover = doc.add_paragraph()
    cover.paragraph_format.space_before = Pt(72)
    r = cover.add_run(eyebrow)
    set_run_font(r, size=12, bold=True, color=TEAL)

    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(8)
    r = h.add_run(title)
    set_run_font(r, size=32, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    r = sub.add_run(subtitle)
    set_run_font(r, size=14, color=MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(28)
    r = meta.add_run(
        "Scale target: 1 million simulations a day, 1 million calls a day,\n"
        "and 1 million live traces a day, at the same time.\n\n"
        "Engineering discussion  ·  26 October 2026\n"
        "Read from the code. Not a load test. Numbers are per replica\n"
        "unless the section says otherwise."
    )
    set_run_font(r, size=12, color=BODY)
    doc.add_page_break()

    in_code = False
    code_buf = []
    para_buf = []
    i = 0

    def flush_para():
        nonlocal para_buf
        text = " ".join(s.strip() for s in para_buf).strip()
        para_buf = []
        if text:
            body_paragraph(doc, text)

    while i < len(lines):
        line = lines[i]
        if in_code:
            if line.strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.3)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_buf))
                set_run_font(run, name="Consolas", size=9, color=NAVY)
                code_buf = []
                in_code = False
            else:
                code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith("```"):
            flush_para()
            in_code = True
            i += 1
            continue
        if line.strip() == "---":
            flush_para()
            i += 1
            continue
        if line.startswith("|"):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_separator(lines[i]):
                    rows.append(parse_table_row(lines[i]))
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith("#"):
            flush_para()
            level = min(len(line) - len(line.lstrip("#")), 3)
            text = line[level:].strip() if False else line.lstrip("#").strip()
            p = doc.add_heading(text, level=level)
            style_heading(p, level)
            i += 1
            continue
        bullet = re.match(r"^(\s*)[-*] (.*)$", line)
        if bullet:
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, "•  " + bullet.group(2).strip())
            i += 1
            continue
        numbered = re.match(r"^(\s*)(\d+)\. (.*)$", line)
        if numbered:
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, f"{numbered.group(2)}.  {numbered.group(3).strip()}")
            i += 1
            continue
        if not line.strip():
            flush_para()
            i += 1
            continue
        para_buf.append(line)
        i += 1
    flush_para()
    doc.save(out_path)
    print(out_path.name, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables))


def slice_lines(start, end):
    """1-based inclusive line numbers."""
    return SRC[start - 1 : end]


part_a = slice_lines(1, 948)
part_a[0] = "# Part A — What breaks at 1 million a day"
part_a[9] = (
    "This is document 1 of 3. It walks the product from the top of the sidebar to the bottom and says what each area can carry today. "
    "**Document 2** is the code shape. **Document 3** is the design-pattern scorecard. "
    "A later note in the full markdown (Part D) covers retention, lost updates, deploys, and spend. It is not in these three files."
)

part_b = [
    "This is document 2 of 3. The scale target is still **1 million simulations, 1 million calls, and 1 million traces a day**. "
    "Document 1 says where the product runs out of room. This document says whether the code is doing extra work on every request, and whether the next feature has a file to go in. "
    "Adding servers does not fix a request that already does the same work twice.",
    "",
] + slice_lines(952, 1223)

part_c = [
    "This is document 3 of 3. The scale target is still **1 million simulations, 1 million calls, and 1 million traces a day**. "
    "A pattern is worth following when a new provider, a new metric, or a new teammate would otherwise edit a giant file and break the last feature. "
    "Document 1 is capacity. Document 2 is the file-by-file code shape. This document is the scorecard.",
    "",
] + slice_lines(1227, 1379)

render(
    part_a,
    ROOT / "01-Capacity-1-million-a-day.docx",
    "DOCUMENT 1 OF 3  ·  PART A",
    "What breaks at\n1 million a day",
    "Sidebar, top to bottom.\nWhat each area can carry today, what is wrong,\nand what to change before adding more servers.",
    "PART A  ·  CAPACITY",
)
render(
    part_b,
    ROOT / "02-Code-shape.docx",
    "DOCUMENT 2 OF 3  ·  PART B",
    "Code shape for\n1 million a day",
    "Extra queries, extra downloads, and files\nthat are thousands of lines long.\nWhere the next feature should go.",
    "PART B  ·  CODE SHAPE",
)
render(
    part_c,
    ROOT / "03-Design-patterns.docx",
    "DOCUMENT 3 OF 3  ·  PART C",
    "Design patterns\nwe do and don't follow",
    "Low-level design, scored against this repo.\nWhat is already a seam, what is only half done,\nand the ten patterns a change should use.",
    "PART C  ·  PATTERNS",
)
